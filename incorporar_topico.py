# -*- coding: utf-8 -*-
import copy
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC = r"C:\Users\alden\OneDrive\Carlos Nina\Representacao_Etico_Disciplinar_OAB-MA.docx"
OUT = r"C:\Users\alden\OneDrive\Carlos Nina\Representacao_Etico_Disciplinar_OAB-MA_atualizada.docx"

d = Document(SRC)

# ---- localizar referências ----
ref_indiv = None
ref_lastdoc = None
for p in d.paragraphs:
    t = p.text
    if ref_indiv is None and 'ndividualiza' in t and 'condutas' in t:
        ref_indiv = p
    if 'Id 164539246' in t and p.style.name == 'List Paragraph':
        ref_lastdoc = p
assert ref_indiv is not None, "não achei a seção Individualização"
assert ref_lastdoc is not None, "não achei o último item de anexos"

# ===== helpers de inserção (antes de ref_indiv) =====
def _set_font(run):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)

def ins_blank():
    p = ref_indiv.insert_paragraph_before()
    return p

def ins_title(text):
    p = ref_indiv.insert_paragraph_before()
    p.alignment = None
    r = p.add_run(text)
    _set_font(r)
    r.font.small_caps = True
    return p

def ins_body(segments, indent_left=None):
    """segments: list of (text, bold, smallcaps)"""
    p = ref_indiv.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent_left is not None:
        p.paragraph_format.left_indent = Cm(indent_left)
    for text, bold, sc in segments:
        r = p.add_run(text)
        _set_font(r)
        r.bold = bold
        r.font.small_caps = sc
    return p

def ins_text(text):
    return ins_body([(text, False, False)])

def ins_quote(text):
    return ins_body([(text, False, False)], indent_left=2.0)

# ===== conteúdo do novo tópico III-A =====
ins_blank()
ins_title("Da Reprovação Judicial das Condutas e do Desmentido Documental das Insinuações")

ins_text("Os elementos a seguir, extraídos dos próprios autos, confirmam que as condutas dos "
         "Representados já foram repelidas pelo Poder Judiciário e desmentem, documentalmente, as "
         "insinuações de que os Representantes se beneficiariam de “tráfico de influência” ou de "
         "“relações de amizade” com magistrados (Fatos III.6 e III.9).")

ins_body([("Das declarações de suspeição por foro íntimo. ", True, False),
          ("Todas as declarações de suspeição ocorreram no Processo 0819450-39.2020 e fundaram-se "
           "exclusivamente em motivo de foro íntimo (art. 145, § 1º, do CPC) — nenhuma reconheceu "
           "parentesco ou amizade com os Representantes, o que desmente as imputações de influência "
           "indevida. Afastaram-se os seguintes magistrados:", False, False)])

juizes = [
    ("1. ", "Maria do Socorro Mendonça Carneiro", " (5ª Vara de Família) — foro íntimo — Id 58239138, pág. 2 — 15/12/2021;"),
    ("2. ", "Lidiane Melo de Souza", " (Auxiliar de Entrância Final) — foro íntimo — Id 58421517 — 17/12/2021;"),
    ("3. ", "Jesus Guanaré de Sousa Borges", " (6ª Vara de Família) — foro íntimo — Id 60349164, pág. 1 — 27/01/2022;"),
    ("4. ", "Ailton Castro Aires", " (1ª Vara de Família) — foro íntimo — Id 62457717, pág. 1 — 11/03/2022;"),
    ("5. ", "Joseane de Jesus Corrêa Bezerra", " (3ª Vara de Família) — foro íntimo — Id 92128344, pág. 1 — 10/05/2023;"),
    ("6. ", "Edilza Barros Ferreira Lopes Viégas", " (Auxiliar de Entrância Final) — foro íntimo — Id 158972913, pág. 1 — 01/09/2025;"),
    ("7. ", "José Augusto Sá Costa Leite", " (Auxiliar de Entrância Final) — foro íntimo — Id 159415861, pág. 1 — 05/09/2025."),
]
for num, nome, resto in juizes:
    ins_body([(num, False, False), (nome, False, True), (resto, False, False)], indent_left=1.0)

ins_text("Acresce-se a Desembargadora Ângela Maria Moraes Salazar, que também declarou suspeição por "
         "foro íntimo ao ser sorteada para o Agravo de Instrumento 0819158-23.2021.8.10.0000 "
         "(transcrito no Id 164539246, pág. 9). A inexistência de qualquer afastamento por "
         "parentesco ou amizade evidencia que os Representados converteram recusas neutras de foro "
         "íntimo em insinuação caluniosa de influência sobre o Judiciário.")

ins_body([("Das decisões que repeliram a conduta dos Representados. ", True, False),
          ("O próprio Juízo da causa já qualificou a conduta dos Representados como tumulto "
           "processual, advertiu quanto à litigância de má-fé e reconheceu o risco de dano à "
           "reputação dos profissionais atacados:", False, False)])

ins_text("(i) Decisão da Juíza Luciany Cristina de Sousa Ferreira, de 17/12/2024 (Id 137187167, "
         "págs. 1/2):")
ins_quote("“...a simples alegação, sem elementos probatórios robustos e concretos, não é suficiente "
          "para fundamentar a remessa de peças processuais à OAB e ao Ministério Público. (...) "
          "não cabe ao Judiciário determinar o encaminhamento sem respaldo probatório mínimo, "
          "evitando-se decisões temerárias que possam causar danos à imagem e à reputação dos "
          "profissionais envolvidos.”")
ins_quote("“...as repetidas declarações de suspeição apresentadas pela parte Ré sugerem, mais que "
          "um legítimo questionamento, uma tentativa deliberada de tumulto processual (...). A "
          "partir de agora, tais questionamentos não serão mais tolerados, sob pena de configuração "
          "de litigância de má-fé (art. 80 do CPC) (...). Este Juízo não admitirá que o processo "
          "seja transformado em um cenário de disputa frívola ou de instrumentalização indevida. O "
          "Direito é ferramenta de pacificação social, e não arena para contendas dilatórias.”")

ins_text("(ii) Decisão da Juíza Nirvana Maria Mourão Barroso, de 12/10/2025 (Id 162760762; cópia "
         "Id 163691546, págs. 7/8):")
ins_quote("“REITERO que não serão admitidos questionamentos infundados acerca da conduta dos "
          "magistrados que atuaram no presente feito, consoante já advertido na decisão de ID "
          "137187167; ADVIRTO as partes de que o uso reiterado de expedientes protelatórios poderá "
          "ensejar a aplicação de multa por litigância de má-fé, nos termos do art. 81 do CPC.”")

ins_text("(iii) Sentença de 09/06/2022 (Id 68685015), no Processo 0842458-45.2020, que revogou as "
         "medidas protetivas por reconhecer que os fatos decorriam de divergências patrimoniais do "
         "divórcio, a serem dirimidas no juízo cível competente.")

ins_text("Tais pronunciamentos, emanados do próprio Poder Judiciário, constituem prova "
         "pré-constituída da improcedência das acusações e do caráter abusivo da conduta dos "
         "Representados, reforçando a justa causa desta representação.")

# ===== anexar novos documentos ao rol (clonando o último item) =====
# 1) trocar o ponto final do último item por ponto e vírgula
for t in ref_lastdoc._p.findall('.//' + qn('w:t')):
    if t.text and '164539246' in t.text:
        t.text = t.text.replace(').', ');')

novos_docs = [
    "Despachos de declaração de suspeição por foro íntimo (Ids 58239138, 58421517, 60349164, 62457717, 92128344, 158972913 e 159415861);",
    "Decisão da Juíza Luciany Cristina de Sousa Ferreira, de 17/12/2024 (Id 137187167);",
    "Decisão da Juíza Nirvana Maria Mourão Barroso, de 12/10/2025 (Id 162760762; cópia Id 163691546);",
    "Sentença de revogação das medidas protetivas, de 09/06/2022, no Processo 0842458-45.2020 (Id 68685015).",
]
anchor = ref_lastdoc._p
for texto in novos_docs:
    new = copy.deepcopy(ref_lastdoc._p)
    ts = new.findall('.//' + qn('w:t'))
    if ts:
        ts[0].text = texto
        for extra in ts[1:]:
            extra.text = ''
    anchor.addnext(new)
    anchor = new

d.save(OUT)
print("Salvo em:", OUT)
print("Total de parágrafos:", len(Document(OUT).paragraphs))
