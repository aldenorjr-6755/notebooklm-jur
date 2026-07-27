# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml.ns import qn

SRC = r"C:\Users\alden\OneDrive\Carlos Nina\Representacao_Etico_Disciplinar_OAB-MA_v3.docx"
OUT = r"C:\Users\alden\OneDrive\Carlos Nina\Representacao_Etico_Disciplinar_OAB-MA_v4.docx"

d = Document(SRC)

def set_single_text(p_el, text):
    runs = p_el.findall(qn('w:r'))
    if not runs:
        return
    first = runs[0]
    for r in runs[1:]:
        p_el.remove(r)
    ts = first.findall(qn('w:t'))
    if ts:
        ts[0].text = text
        ts[0].set(qn('xml:space'), 'preserve')
        for extra in ts[1:]:
            first.remove(extra)

def norm(s):
    return s.strip().lower()

to_delete = []
for p in d.paragraphs:
    t = norm(p.text)
    if t == '(cumulada com pedido de desagravo público)':
        to_delete.append(p)
    elif t.startswith('direito ao desagravo público'):
        to_delete.append(p)
    elif t.startswith('a concessão do desagravo público'):
        to_delete.append(p)
    elif t == 'rol de testemunhas':
        to_delete.append(p)
    elif t.startswith('[a informar') and 'art. 57' in t:
        to_delete.append(p)
    elif t.startswith('a produção de todas as provas'):
        set_single_text(p._p, ("a produção de todas as provas admitidas, especialmente: "
                                "(a) documental (cópias das peças identificadas por Id e página "
                                "— anexos); e (b) depoimento pessoal dos Representados;"))

for p in to_delete:
    p._p.getparent().remove(p._p)

d.save(OUT)

# verificação
d2 = Document(OUT)
print("Salvo em:", OUT)
print("Parágrafos:", len(d2.paragraphs))
hits = [p.text for p in d2.paragraphs if 'desagravo' in p.text.lower() or 'testemunh' in p.text.lower()]
print("Resíduos de desagravo/testemunha:", hits if hits else "NENHUM")
for p in d2.paragraphs:
    if p.text.strip().startswith('a produção de todas as provas'):
        print("Pedido de provas ->", p.text)
