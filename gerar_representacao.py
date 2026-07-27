# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
import re

doc = Document()

# Page setup (A4, margens ABNT-ish)
sec = doc.sections[0]
sec.page_height = Cm(29.7)
sec.page_width = Cm(21.0)
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(2.0)

# Base style
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
pf = style.paragraph_format
pf.line_spacing = 1.5
pf.space_after = Pt(0)

def add(text="", *, bold=False, align='just', italic=False, first_indent=True,
        space_before=0, space_after=6, size=12, all_caps=False):
    p = doc.add_paragraph()
    a = {'just': WD_ALIGN_PARAGRAPH.JUSTIFY, 'center': WD_ALIGN_PARAGRAPH.CENTER,
         'left': WD_ALIGN_PARAGRAPH.LEFT, 'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.alignment = a
    pfmt = p.paragraph_format
    pfmt.space_before = Pt(space_before)
    pfmt.space_after = Pt(space_after)
    if first_indent and align == 'just':
        pfmt.first_line_indent = Cm(1.25)
    if text:
        run = p.add_run(text.upper() if all_caps else text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
    return p

def add_rich(segments, *, align='just', first_indent=True, space_after=6, space_before=0):
    """segments: list of (text, bold, italic)"""
    p = doc.add_paragraph()
    p.alignment = {'just': WD_ALIGN_PARAGRAPH.JUSTIFY, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'left': WD_ALIGN_PARAGRAPH.LEFT}[align]
    pfmt = p.paragraph_format
    pfmt.space_before = Pt(space_before)
    pfmt.space_after = Pt(space_after)
    if first_indent and align == 'just':
        pfmt.first_line_indent = Cm(1.25)
    for text, bold, italic in segments:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
    return p

def heading(text):
    add(text, bold=True, align='left', first_indent=False, space_before=12, space_after=6)

def subheading(text):
    add(text, bold=True, align='just', first_indent=False, space_before=8, space_after=4)

# ===== ENDEREÇAMENTO =====
add("EXCELENTÍSSIMO SENHOR PRESIDENTE DO TRIBUNAL DE ÉTICA E DISCIPLINA DO CONSELHO SECCIONAL DA ORDEM DOS ADVOGADOS DO BRASIL – SECCIONAL DO MARANHÃO (OAB/MA)",
    bold=True, align='just', first_indent=False, space_after=18)

for _ in range(6):
    add("", first_indent=False, space_after=0)

# ===== QUALIFICAÇÃO ATIVA =====
add_rich([
    ("CARLOS SEBASTIÃO SILVA NINA", True, False),
    (", advogado, casado, inscrito na OAB/MA sob o nº 4.870-A e na OAB/SP sob o nº 151.986, "
     "portador do CPF nº 012.181.313-49, e ", False, False),
    ("ENIDE MARIA AQUINO NINA", True, False),
    (", advogada, casada com o primeiro Representante, inscrita na OAB/MA sob o nº 5.397, "
     "portadora do CPF nº 428.420.933-72, ambos com endereço profissional comum na Avenida dos "
     "Holandeses, Quadra 05, Lote 02, Sala 811-A, Edifício Marcus Barbosa Intelligent Office, "
     "Calhau, São Luís/MA, CEP 65.071-380, endereço eletrônico carlos.nina@yahoo.com.br, onde "
     "recebem intimações, vêm, respeitosamente, à presença de Vossa Excelência, com fundamento "
     "nos arts. 70 e 72 da Lei nº 8.906/94 (EOAB) e no art. 55 e seguintes do Código de Ética e "
     "Disciplina da OAB (Resolução CFOAB nº 02/2015), apresentar a presente", False, False),
], space_after=12)

add("REPRESENTAÇÃO ÉTICO-DISCIPLINAR", bold=True, align='center', first_indent=False, space_after=0)
add("(cumulada com pedido de DESAGRAVO PÚBLICO)", bold=True, align='center', first_indent=False, space_after=12)

add("em face de:", first_indent=False, space_after=6)

reus = [
    ("1) MARIA DA GLÓRIA COSTA GONÇALVES DE SOUSA AQUINO",
     ", advogada, inscrita na OAB/MA sob o nº 6.399-A e na OAB/RJ sob o nº 105.640, "
     "CPF nº 042.775.647-20, com endereço na Avenida Presidente Juscelino Kubitschek, Quadra 16/17, "
     "Edifício Quinta da Boa Vista, Apto 201, Quintas do Calhau, São Luís/MA, CEP 65.072-005;"),
    ("2) EDUARDO ALEXANDRE COSTA CORRÊA",
     ", advogado, inscrito na OAB/MA sob o nº 5.211-A, CPF nº 557.060.533-91, com endereço na "
     "Avenida Senador Vitorino Freire, nº 1.958, salas 117/118, Edifício Business Center, Areinha, "
     "São Luís/MA, CEP 65.030-015;"),
    ("3) MILTON RICARDO LUSO CALADO",
     ", advogado, inscrito na OAB/MA sob o nº 5.108, com CPF a ser complementado pela Secretaria "
     "deste E. Tribunal de Ética e Disciplina mediante consulta aos assentamentos cadastrais da "
     "OAB/MA, com endereço profissional na Avenida Senador Vitorino Freire, nº 1.958, salas 117/118, "
     "Edifício Business Center, Areinha, São Luís/MA, CEP 65.030-015, endereço eletrônico "
     "mrlcalado@ig.com.br;"),
    ("4) THYENES DE OLIVEIRA CHAGAS CORRÊA",
     ", advogada, inscrita na OAB/MA sob o nº 5.114, com CPF a ser complementado pela Secretaria "
     "deste E. Tribunal de Ética e Disciplina mediante consulta aos assentamentos cadastrais da "
     "OAB/MA, com o mesmo endereço profissional acima, endereço eletrônico thyenes@gmail.com;"),
    ("5) IRANILDE TEIXEIRA DE JESUS ANDRADE",
     ", advogada, inscrita na OAB/MA sob o nº 20.465, com CPF a ser complementado pela Secretaria "
     "deste E. Tribunal de Ética e Disciplina mediante consulta aos assentamentos cadastrais da "
     "OAB/MA, com o mesmo endereço profissional acima,"),
]
for nome, resto in reus:
    add_rich([(nome, True, False), (resto, False, False)], space_after=6)

add("pelos fatos e fundamentos a seguir expostos.", first_indent=False, space_after=6)

# ===== I COMPETENCIA =====
heading("I – DA COMPETÊNCIA E DA LEGITIMIDADE")
add_rich([
    ("A competência é do ", False, False),
    ("Tribunal de Ética e Disciplina do Conselho Seccional da OAB/MA", True, False),
    (", pois todas as condutas foram praticadas em petições protocoladas em processos que tramitam "
     "na Comarca de São Luís/MA (“o poder de punir compete exclusivamente ao Conselho Seccional "
     "em cuja base tenha ocorrido a infração” — Manual de Procedimentos, Segunda Parte, item 25). "
     "Os Representantes têm legitimidade ativa, como advogados ofendidos no exercício profissional; "
     "os Representados, regularmente inscritos na OAB/MA, sujeitam-se ao poder disciplinar desta Casa "
     "(art. 70 do EOAB).", False, False),
])

# ===== II ORIGEM =====
heading("II – DA ORIGEM DA CONTROVÉRSIA")
add("Os Representantes atuam como advogados constituídos do Sr. Heráclito Aquino Junior em três "
    "processos na Comarca de São Luís/MA, decorrentes da dissolução do vínculo conjugal entre seu "
    "cliente e a Representada Maria da Glória. Registre-se, desde logo, que os Representantes são "
    "casados entre si e mantêm relação de parentesco com o cliente: a Representante Enide Maria "
    "Aquino Nina é irmã do Sr. Heráclito Aquino Junior, e o Representante Carlos Sebastião Silva "
    "Nina é seu cunhado. Tal vínculo familiar — legítimo e que em nada obsta o exercício do mandato "
    "— foi, contudo, deliberadamente explorado pelos Representados para, a pretexto de “relação "
    "de parentesco”, construir as insinuações de “tráfico de influência”, “conluio "
    "familiar” e atuação concertada para “frustrar a meação”, como adiante se "
    "demonstrará. A condição de irmã e cunhado do requerido, por si só, jamais poderia servir de "
    "fundamento à imputação de crimes e desvios de conduta, sendo essa exploração espúria do "
    "parentesco um dos eixos centrais das ofensas ora representadas.")
add("Os Representados Eduardo Alexandre Costa Corrêa, Milton Ricardo Luso Calado, Thyenes de "
    "Oliveira Chagas Corrêa e Iranilde Teixeira de Jesus Andrade integram a banca Calado & Corrêa "
    "Advogados Associados, patrona da Sra. Maria da Glória, a qual, a partir de 27/11/2023, passou "
    "a subscrever pessoalmente as petições, em causa própria. Ao longo de mais de quatro anos, os "
    "Representados dirigiram aos Representantes uma sucessão crescente de ofensas, insinuações "
    "injuriosas, imputações difamatórias e acusações caluniosas de crimes.")

# ===== III FATOS =====
heading("III – DOS FATOS E DAS INFRAÇÕES (LINHA DO TEMPO E AUTORIA)")

def fato(titulo, ref, subs, corpo_segments, regra):
    subheading(titulo)
    add_rich([(ref, False, True)], first_indent=False, space_after=2)
    add_rich([("Subscritores: ", True, False), (subs, True, False)], first_indent=False, space_after=4)
    add_rich(corpo_segments, space_after=4)
    add_rich([("Regra infringida: ", True, False), (regra, False, False)],
             first_indent=False, space_after=8)

fato("III.1 – 17/03/2021 — Imputação de intenção maliciosa de prejudicar",
     "Processos nº 0842458-45.2020 (Id 42767466, pág. 3) e nº 0800032-81.2021 (Id 42768391, pág. 3).",
     "EDUARDO CORRÊA, MILTON CALADO, THYENES CORRÊA e IRANILDE ANDRADE.",
     [("“...omissão que evidencia ainda mais a intenção de causar prejuízos à Denunciante, que "
       "já se encontra em situação de vulnerabilidade”", False, True),
      (" (referindo-se aos patronos).", False, False)],
     "arts. 27 e 28 do CED.")

fato("III.2 – 28/04/2021 — Imputação de intenção de difamar perante OAB e CGJ",
     "Processos nº 0842458-45.2020 (Id 44801025, pág. 2) e nº 0800032-81.2021 (Id 44801018, pág. 2).",
     "EDUARDO CORRÊA, MILTON CALADO, THYENES CORRÊA e IRANILDE ANDRADE.",
     [("“...os fatos foram expostos pelos advogados do Denunciado em ofícios encaminhados à "
       "OAB/MA e à CGJ, com a evidente intenção de expô-la a situação que venha a atingir sua "
       "reputação (...) que não o de causar constrangimento à Denunciante.”", False, True)],
     "arts. 27 e 28 do CED; art. 2º, parágrafo único, II, do CED.")

fato("III.3 – 14/12/2021 — Imputação de má-fé e “tentativa espúria” de ludibriar o Juízo",
     "Processo nº 0842458-45.2020 (Id 58205286, pág. 2).",
     "EDUARDO CORRÊA, MILTON CALADO, THYENES CORRÊA e IRANILDE ANDRADE.",
     [("“...MERA BOA-FÉ APARENTE, com o propósito de ludibriar esse r. Juízo, na tentativa "
       "espúria de convencê-lo de que a violência patrimonial não está a ocorrer!”", False, True)],
     "arts. 27 e 28 do CED; art. 31 do EOAB.")

fato("III.4 – 09/08/2022 — Imputação de “perseguição”, “devassa”, “obsessão” e “conduta vexatória”",
     "Processo nº 0842458-45.2020 (Id 73383832, págs. 4 e 9) e reiteração no Processo nº 0819450-39.2020 (Id 73383832, págs. 4–5).",
     "EDUARDO CORRÊA, MILTON CALADO, THYENES CORRÊA e IRANILDE ANDRADE.",
     [("“A embargante vem sofrendo perseguição por parte dos advogados do embargado que vêm "
       "promovendo uma devassa (...). Ou seja, tais fatos demonstram que a obsessão existe por parte "
       "dos advogados (cunhado e irmã) do embargado;”", False, True),
      (" e ", False, False),
      ("“...Uma conduta vexatória que fere o Código de Ética e Disciplina da OAB...”", False, True)],
     "art. 27, caput e § 2º, do CED; art. 28 do CED.")

fato("III.5 – 27/11/2023 — Linguagem injuriosa: “estapafúrdio” e “pilhéria”",
     "Processo nº 0819450-39.2020 (Id 107341869, págs. 19/20).",
     "EDUARDO CORRÊA, MILTON CALADO, THYENES CORRÊA e MARIA DA GLÓRIA (em causa própria). Iranilde não subscreveu.",
     [("“...os argumentos trazidos na contestação são tão estapafúrdios que fazem verdadeira "
       "pilhéria com a Justiça!”", False, True)],
     "arts. 27 e 28 do CED.")

fato("III.6 – 18/09/2024 — Imputação de “tráfico de influência”, “prática ilícita/coautoria” e fraude",
     "Processo nº 0819450-39.2020 (Id 129629168, págs. 4, 11 e 24/25).",
     "EDUARDO CORRÊA e MARIA DA GLÓRIA (em causa própria).",
     [("(pág. 4) ", False, False),
      ("“...eventual tráfico de influência na tentativa espúria de obstaculizar o acesso à "
       "justiça...”", False, True),
      ("; (pág. 11) ", False, False),
      ("“...a prática ilícita vem contando com o apoio e a coautoria (...) beneficiando de modo "
       "indireto os patronos do Requerido...”", False, True),
      ("; (pág. 24/25) ", False, False),
      ("“...manipulação fraudulenta dos rendimentos (...) com o (...) objetivo de se locupletarem "
       "(...) caracterizando litigância de má-fé...”", False, True)],
     "art. 34, XV, do EOAB (imputação de crime); arts. 27 e 28 do CED; art. 2º, parágrafo único, do CED.")

fato("III.7 – 28/01/2025 — Pedido de representação ao MP por “prática de crimes”",
     "Processo nº 0819450-39.2020 (Id 139598410, págs. 19/20).",
     "EDUARDO CORRÊA e MARIA DA GLÓRIA (em causa própria).",
     [("“...Ministério Público, para averiguação de eventual prática de crimes aos advogados: "
       "(...) Carlos Sebastião Silva Nina (...); Enide Maria Aquino Nina”", False, True)],
     "art. 34, XV, do EOAB; art. 27 do CED; art. 31 do EOAB.")

fato("III.8 – 12/08/2025 — Imputação de estelionato, lavagem de dinheiro, sonegação, conluio e locupletamento; pedido de quebra de sigilo dos advogados",
     "Processo nº 0819450-39.2020 (Id 157089391, págs. 2, 11/12, 12/13, 14/15, 15/16 e 27/28).",
     "EDUARDO CORRÊA e MARIA DA GLÓRIA (em causa própria).",
     [("(pág. 2) ", False, False),
      ("“...manobras (...) às escondidas e em conluio...”", False, True),
      ("; (pág. 11/12) ", False, False),
      ("“...os advogados constituídos nos autos estão auxiliando na prática da sonegação, da "
       "dilapidação rápida dos bens comuns...”", False, True),
      ("; (pág. 12/13) ", False, False),
      ("“...os advogados – irmã e cunhado do Requerido, estão se beneficiando (...) "
       "locupletando-se indevidamente...”", False, True),
      ("; (pág. 14/15) pedido de ", False, False),
      ("“quebra do sigilo fiscal, bancário e de cartões de crédito de: Enide Maria Aquino Nina "
       "(...) Carlos Sebastião Silva Nina”", False, True),
      ("; (pág. 15/16) ", False, False),
      ("“...averiguação de eventual prática de crime de estelionato (artigo 171, do Código Penal) "
       "aos advogados...”", False, True),
      (" e ", False, False),
      ("“branqueamento de capitais (lavagem de dinheiro)”", False, True),
      (".", False, False)],
     "art. 34, XV, do EOAB; art. 34, XXV, c/c § 1º, “b”; arts. 27 e 28 do CED; art. 2º, parágrafo único, do CED.")

fato("III.9 – 14/08/2025 — Imputação de denegrição da imagem, “stalking” e insinuação de influência sobre o Judiciário",
     "Processo nº 0819450-39.2020 (Id 157357708, pág. 4).",
     "EDUARDO CORRÊA e MARIA DA GLÓRIA (em causa própria).",
     [("“...denegrição da imagem e da reputação pessoal e profissional da Requerente pelos "
       "advogados do Requerido”", False, True),
      ("; ", False, False),
      ("“...perseguição insistente (stalking) de terceiros - pessoas vinculadas aos advogados "
       "do Requerido - nas redes sociais”", False, True),
      ("; ", False, False),
      ("“...suspeições de juízes, provenientes das ‘relações de amizade com os advogados "
       "do Requerido’”", False, True),
      (".", False, False)],
     "arts. 27 e 28 do CED; art. 2º, parágrafo único, II e III, do CED.")

fato("III.10 – 30/10/2025 — Reiteração das imputações de locupletamento, lavagem e sonegação",
     "Processo nº 0819450-39.2020 (Id 164539246, págs. 33, 34 e 34/35).",
     "EDUARDO CORRÊA e MARIA DA GLÓRIA (em causa própria).",
     [("Reiteração do pedido de quebra de sigilo dos Representantes (pág. 33), do ofício à OAB/MA "
       "por “aplicação ilícita ou desonesta” e “locupletação ilícita” (pág. 34) e "
       "da imputação de ", False, False),
      ("“sonegação, dilapidação e branqueamento de capitais (lavagem de dinheiro)”", False, True),
      (" (pág. 34/35).", False, False)],
     "art. 34, XV e XXV, do EOAB; arts. 27 e 28 do CED; continuidade/reiteração (art. 37, II, do EOAB).")

# ===== IV INDIVIDUALIZACAO =====
heading("IV – DA INDIVIDUALIZAÇÃO DAS CONDUTAS")
add_rich([("A prova documental revela ", False, False),
          ("graus distintos de responsabilidade", True, False),
          (", que devem orientar a dosimetria (arts. 35 a 37 do EOAB):", False, False)])

add_rich([("IV.1 – EDUARDO ALEXANDRE COSTA CORRÊA (OAB/MA 5.211-A). ", True, False),
          ("Subscritor de TODAS as dez petições (2021–2025), responde pela integralidade das condutas "
           "— das ofensas iniciais às imputações criminais mais graves (tráfico de influência, "
           "estelionato, lavagem de dinheiro, sonegação, locupletamento). Como patrono em todo o "
           "período, tinha domínio técnico sobre o conteúdo das peças. A reiteração por mais de quatro "
           "anos evidencia conduta deliberada e contumaz, atraindo o art. 34, XV e XXV, do EOAB, e a "
           "possibilidade de suspensão (art. 37, I e II).", False, False)])

add_rich([("IV.2 – MARIA DA GLÓRIA COSTA GONÇALVES DE SOUSA AQUINO (OAB/MA 6.399-A). ", True, False),
          ("Subscreveu seis petições a partir de 27/11/2023, em causa própria (parte e advogada), "
           "justamente as de maior gravidade (linguagem injuriosa, tráfico de influência, fraude, "
           "estelionato, lavagem de dinheiro, sonegação, locupletamento). Sua condição de advogada "
           "agrava a reprovabilidade, pois conhecia os limites dos arts. 27 e 28 do CED e do art. 34, "
           "XV, do EOAB. Sujeita-se aos arts. 36 e 37 do EOAB.", False, False)])

add_rich([("IV.3 – MILTON RICARDO LUSO CALADO (OAB/MA 5.108) e THYENES DE OLIVEIRA CHAGAS CORRÊA "
           "(OAB/MA 5.114). ", True, False),
          ("Subscreveram as cinco primeiras manifestações infracionais (Fatos III.1 a III.5, de "
           "17/03/2021 a 27/11/2023). Sua responsabilidade restringe-se às ofensas de urbanidade e "
           "linguagem (arts. 27 e 28 do CED) e ao decoro (art. 31 do EOAB) — não subscreveram nenhuma "
           "das imputações de crime (que se iniciam em 18/09/2024). Sujeitam-se, em regra, à pena de "
           "censura (art. 36, II, do EOAB).", False, False)])

add_rich([("IV.4 – IRANILDE TEIXEIRA DE JESUS ANDRADE (OAB/MA 20.465). ", True, False),
          ("Subscreveu as quatro primeiras manifestações (Fatos III.1 a III.4, de 17/03/2021 a "
           "09/08/2022), deixando de figurar a partir de 27/11/2023. Sua responsabilidade é a mais "
           "restrita, limitada às ofensas de urbanidade e linguagem dos Fatos III.1 a III.4 (arts. 27 "
           "e 28 do CED), sem qualquer participação nas imputações criminais. Sujeita-se à pena de "
           "censura (art. 36, II, do EOAB).", False, False)])

# ===== V DIREITO =====
heading("V – DO DIREITO")
add_rich([("V.1 – Do dever de urbanidade e da linguagem polida (arts. 27 e 28 do CED). ", True, False),
          ("O art. 27 impõe tratar os colegas “com respeito e consideração”, e o § 2º "
           "determina, havendo “ofensa à honra do advogado”, a instauração de “processo "
           "ético-disciplinar”. O art. 28 erige a “linguagem escorreita e polida” a "
           "imperativo da correta atuação.", False, False)])
add_rich([("V.2 – Do decoro (art. 31 do EOAB) e dos deveres do art. 2º do CED. ", True, False),
          ("Exige-se conduta que contribua “para o prestígio da classe e da advocacia” e "
           "atuação com “honestidade, decoro, veracidade, lealdade, dignidade e boa-fé”.", False, False)])
add_rich([("V.3 – Da imputação de crime a colega (art. 34, XV, do EOAB). ", True, False),
          ("Constitui infração “fazer (...) imputação a terceiro de fato definido como "
           "crime”. Os Representados Eduardo e Maria da Glória imputaram nominalmente aos "
           "Representantes estelionato (art. 171, CP), lavagem de dinheiro (Lei nº 9.613/98), sonegação "
           "e tráfico de influência, sem lastro probatório, requerendo até a quebra de seus sigilos.", False, False)])
add_rich([("V.4 – Da conduta incompatível (art. 34, XXV, c/c § 1º, “b”, do EOAB). ", True, False),
          ("A reiteração de acusações criminais graves e infundadas configura conduta incompatível, "
           "podendo caracterizar “incontinência pública e escandalosa”.", False, False)])
add_rich([("V.5 – Das sanções (arts. 35 a 37 do EOAB). ", True, False),
          ("A violação a preceito do CED enseja censura (art. 36, II); a infração ao art. 34, XXV, e a "
           "reincidência sujeitam à suspensão (art. 37, I e II).", False, False)])
add_rich([("V.6 – Do direito ao desagravo público (art. 18 do Regulamento Geral). ", True, False),
          ("“O inscrito na OAB, quando ofendido comprovadamente em razão do exercício "
           "profissional (...), tem direito ao desagravo público promovido pelo Conselho competente "
           "(...).” As ofensas foram dirigidas aos Representantes exclusivamente em razão do "
           "exercício da advocacia, e não por motivo pessoal, doutrinário, político ou religioso "
           "(afastado o arquivamento do § 3º). O desagravo, nos termos do § 9º, “não depende de "
           "concordância do ofendido (...), devendo ser promovido a critério do Conselho”.", False, False)])

# ===== VI PRESCRICAO =====
heading("VI – DA TEMPESTIVIDADE (PRESCRIÇÃO)")
add("Não há prescrição. A prescrição quinquenal (art. 43, caput, do EOAB) conta-se da constatação "
    "oficial do fato pela OAB. Como as condutas se estenderam até 30/10/2025 e esta representação é "
    "apresentada nesta data, o poder punitivo está integralmente preservado, inclusive quanto aos "
    "fatos iniciais, em razão da continuidade infracional.")

# ===== VII PEDIDOS =====
heading("VII – DOS PEDIDOS")
add("Ante o exposto, requerem os Representantes:", first_indent=False)

def pedido(letra, segments):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pfmt = p.paragraph_format
    pfmt.left_indent = Cm(1.0)
    pfmt.space_after = Pt(6)
    r0 = p.add_run(letra + " ")
    r0.bold = True
    for text, bold, italic in segments:
        r = p.add_run(text); r.bold = bold; r.italic = italic

pedido("a)", [("O recebimento e a autuação da representação, com a notificação dos cinco Representados "
               "para apresentarem defesa prévia no prazo de 15 (quinze) dias úteis;", False, False)])
pedido("a.1)", [("Considerando que os Representados Milton Ricardo Luso Calado, Thyenes de Oliveira "
                 "Chagas Corrêa e Iranilde Teixeira de Jesus Andrade encontram-se plena e "
                 "inequivocamente individualizados por seus números de inscrição na OAB/MA (nºs 5.108, "
                 "5.114 e 20.465, respectivamente), e não constando seus números de CPF dos documentos "
                 "que instruem esta representação, requer-se que a Secretaria deste E. Tribunal de Ética "
                 "e Disciplina complemente tais dados a partir dos assentamentos cadastrais desta "
                 "Seccional, nos termos do art. 57 do Código de Ética e Disciplina, viabilizando a "
                 "regular notificação dos Representados;", False, False)])
pedido("b)", [("A instauração do processo ético-disciplinar em face de EDUARDO ALEXANDRE COSTA CORRÊA "
               "(OAB/MA 5.211-A), MARIA DA GLÓRIA COSTA GONÇALVES DE SOUSA AQUINO (OAB/MA 6.399-A), "
               "MILTON RICARDO LUSO CALADO (OAB/MA 5.108), THYENES DE OLIVEIRA CHAGAS CORRÊA (OAB/MA "
               "5.114) e IRANILDE TEIXEIRA DE JESUS ANDRADE (OAB/MA 20.465);", False, False)])
pedido("c)", [("No mérito, o julgamento de PROCEDÊNCIA, com:", False, False)])

def subpedido(letra, segments):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pfmt = p.paragraph_format
    pfmt.left_indent = Cm(2.0)
    pfmt.space_after = Pt(6)
    r0 = p.add_run(letra + " "); r0.bold = True
    for text, bold, italic in segments:
        r = p.add_run(text); r.bold = bold; r.italic = italic

subpedido("c.1)", [("quanto a EDUARDO ALEXANDRE COSTA CORRÊA, o reconhecimento das infrações aos arts. "
                    "27, 28 e 2º, parágrafo único, do CED, e ao art. 31 e ao art. 34, XV e XXV (c/c § "
                    "1º), da Lei nº 8.906/94, relativamente às dez petições por ele subscritas (Fatos "
                    "III.1 a III.10), com aplicação das sanções dos arts. 36 e 37 do EOAB, observada a "
                    "reiteração/continuidade;", False, False)])
subpedido("c.2)", [("quanto a MARIA DA GLÓRIA COSTA GONÇALVES DE SOUSA AQUINO, o reconhecimento das "
                    "mesmas infrações, relativamente às seis petições por ela subscritas em causa "
                    "própria (Fatos III.5 a III.10), com aplicação dos arts. 36 e 37 do EOAB, "
                    "considerada sua dupla condição de parte e advogada como agravamento;", False, False)])
subpedido("c.3)", [("quanto a MILTON RICARDO LUSO CALADO e THYENES DE OLIVEIRA CHAGAS CORRÊA, o "
                    "reconhecimento das infrações aos arts. 27 e 28 do CED e ao art. 31 do EOAB, "
                    "relativamente às cinco petições por eles subscritas (Fatos III.1 a III.5), com "
                    "aplicação da sanção de censura (art. 36, II, do EOAB);", False, False)])
subpedido("c.4)", [("quanto a IRANILDE TEIXEIRA DE JESUS ANDRADE, o reconhecimento das infrações aos "
                    "arts. 27 e 28 do CED, relativamente às quatro petições por ela subscritas (Fatos "
                    "III.1 a III.4), com aplicação da sanção de censura (art. 36, II, do EOAB);", False, False)])

pedido("d)", [("A concessão do DESAGRAVO PÚBLICO em favor dos Representantes (art. 18 do Regulamento "
               "Geral), com a remessa do pedido à Diretoria/órgão competente do Conselho Seccional "
               "para instrução e designação de sessão de desagravo;", False, False)])
pedido("e)", [("A produção de todas as provas admitidas, especialmente: e.1) documental (cópias das "
               "petições, identificadas por Id e página — docs. anexos); e.2) testemunhal (rol abaixo "
               "— art. 57, III, do CED); e.3) depoimento pessoal dos Representados;", False, False)])
pedido("f)", [("A intimação dos Representantes, no endereço profissional e no e-mail "
               "carlos.nina@yahoo.com.br, de todos os atos do processo.", False, False)])

# ===== VIII PROVAS =====
heading("VIII – DAS PROVAS E DOCUMENTOS ANEXOS")
docs = [
    "Doc. 01 — Comprovantes de inscrição na OAB/MA dos Representantes;",
    "Doc. 02 — Petição de 17/03/2021 (Ids 42767466 e 42768391);",
    "Doc. 03 — Petição de 28/04/2021 (Ids 44801025 e 44801018);",
    "Doc. 04 — Petição de 14/12/2021 (Id 58205286);",
    "Doc. 05 — Manifestação de 09/08/2022 (Id 73383832);",
    "Doc. 06 — Petição de saneamento de 27/11/2023 (Id 107341869);",
    "Doc. 07 — Petição de 18/09/2024 (Id 129629168);",
    "Doc. 08 — Embargos de 28/01/2025 (Id 139598410);",
    "Doc. 09 — Petição intercorrente de 12/08/2025 (Id 157089391);",
    "Doc. 10 — Complemento de 14/08/2025 (Id 157357708);",
    "Doc. 11 — Embargos de 30/10/2025 (Id 164539246).",
]
for d in docs:
    p = doc.add_paragraph(d)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.space_after = Pt(3)

add_rich([("Rol de testemunhas: ", True, False),
          ("[A INFORMAR — até cinco, art. 57, III, do CED].", False, False)],
         first_indent=False, space_before=6)

# ===== FECHO =====
add("Termos em que,", first_indent=False, space_before=14, space_after=0)
add("pede deferimento.", first_indent=False, space_after=14)
add("São Luís/MA, 17 de junho de 2026.", first_indent=False, space_after=36, align='center')

add("_________________________________________", align='center', first_indent=False, space_after=0)
add("CARLOS SEBASTIÃO SILVA NINA", bold=True, align='center', first_indent=False, space_after=0)
add("OAB/MA nº 4.870-A", align='center', first_indent=False, space_after=30)

add("_________________________________________", align='center', first_indent=False, space_after=0)
add("ENIDE MARIA AQUINO NINA", bold=True, align='center', first_indent=False, space_after=0)
add("OAB/MA nº 5.397", align='center', first_indent=False, space_after=0)

out = r"C:\Users\alden\.notebooklm\Representacao_Etico_Disciplinar_OAB-MA.docx"
doc.save(out)
print("Salvo em:", out)
