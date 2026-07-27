# -*- coding: utf-8 -*-
"""Gera a Resposta à Acusação de José Raimundo Sales Chaves Júnior (.docx)."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

doc = Document()

# Margens
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(3.0); s.right_margin = Cm(2.0)

# Estilo base
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
pf = normal.paragraph_format
pf.line_spacing = 1.5
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
pf.space_after = Pt(6)


def p(text="", bold=False, italic=False, align='justify', first_line=True, size=12, space_after=6):
    par = doc.add_paragraph()
    par.paragraph_format.line_spacing = 1.5
    par.paragraph_format.space_after = Pt(space_after)
    amap = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'center': WD_ALIGN_PARAGRAPH.CENTER,
            'right': WD_ALIGN_PARAGRAPH.RIGHT, 'left': WD_ALIGN_PARAGRAPH.LEFT}
    par.alignment = amap[align]
    if first_line and align == 'justify':
        par.paragraph_format.first_line_indent = Cm(1.25)
    run = par.add_run(text)
    run.bold = bold; run.italic = italic
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'
    return par


def rich(segments, align='justify', first_line=True, space_after=6):
    """segments: list of (text, bold, italic)."""
    par = doc.add_paragraph()
    par.paragraph_format.line_spacing = 1.5
    par.paragraph_format.space_after = Pt(space_after)
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if align == 'justify' else WD_ALIGN_PARAGRAPH.CENTER
    if first_line and align == 'justify':
        par.paragraph_format.first_line_indent = Cm(1.25)
    for t, b, i in segments:
        r = par.add_run(t)
        r.bold = b; r.italic = i
        r.font.size = Pt(12); r.font.name = 'Times New Roman'
    return par


def heading(text):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(12)
    par.paragraph_format.space_after = Pt(8)
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = par.add_run(text)
    r.bold = True; r.font.size = Pt(12); r.font.name = 'Times New Roman'
    return par


def quote(text):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(4.0)
    par.paragraph_format.line_spacing = 1.0
    par.paragraph_format.space_after = Pt(8)
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = par.add_run(text)
    r.italic = True; r.font.size = Pt(11); r.font.name = 'Times New Roman'
    return par


# ------------------------------------------------------------------ ENDEREÇAMENTO
p("EXCELENTÍSSIMO SENHOR DOUTOR JUIZ DE DIREITO DA 2ª VARA DO TRIBUNAL DO JÚRI DO TERMO JUDICIÁRIO DE SÃO LUÍS — COMARCA DA ILHA DE SÃO LUÍS/MA",
  bold=True, align='justify', first_line=False, space_after=24)

p("", first_line=False, space_after=0)
p("", first_line=False, space_after=0)

rich([("Autos nº ", True, False), ("0860085-86.2025.8.10.0001", True, False)],
     align='right', first_line=False, space_after=24)

# ------------------------------------------------------------------ PREÂMBULO
rich([
    ("JOSÉ RAIMUNDO SALES CHAVES JÚNIOR", True, False),
    (", já qualificado nos autos do processo-crime em epígrafe, nascido em 08/05/1974, "
     "filho de José Raimundo Sales Chaves e de Maria da Costa Chaves, portador do RG nº "
     "000097378698-1 e inscrito no CPF sob o nº 529.102.103-91, residente e domiciliado na "
     "Alameda Jaú, nº 10, bairro Olho D’Água, São Luís/MA, atualmente em liberdade provisória "
     "(concedida no HC nº 0808210-51.2023.8.10.0000), por seu advogado que esta subscreve, "
     "vem, respeitosamente, à presença de Vossa Excelência, com fundamento no ", False, False),
    ("art. 406, §3º, do Código de Processo Penal", True, False),
    (", apresentar tempestivamente sua", False, False),
], first_line=True)

p("RESPOSTA À ACUSAÇÃO", bold=True, align='center', first_line=False, space_after=12)

p("pelas razões de fato e de direito a seguir expostas.", first_line=True, space_after=18)

# ------------------------------------------------------------------ I - SÍNTESE
heading("I — DA TEMPESTIVIDADE E DA SÍNTESE DA IMPUTAÇÃO")

p("A presente resposta é tempestiva, apresentada no prazo legal a contar da intimação do "
  "subscritor para o ato, em estrito cumprimento ao contraditório e à ampla defesa "
  "(art. 5º, LV, da Constituição Federal).")

rich([
    ("O Ministério Público, em denúncia oferecida em 22/05/2023 e aditada em 16/06/2023, "
     "recebida em 19/06/2023, imputou ao ora respondente a condição de ", False, False),
    ("mandante e coautor", True, False),
    (" do homicídio qualificado consumado de Marcelo Martins Mendes (art. 121, §2º, incisos "
     "II e IV, do CP), do homicídio qualificado tentado contra Felix da Silva Mendes Filho "
     "(art. 121, §2º, incisos II e IV, c/c art. 14, II, do CP) e do crime de associação "
     "criminosa (art. 288 do CP).", False, False),
])

rich([
    ("Segundo a peça acusatória, o respondente teria, motivado por desavença anterior relativa "
     "a uma máquina motoniveladora (“patrol”), ", False, False),
    ("contratado os executores", False, True),
    (" e ", False, False),
    ("cedido um dos veículos de apoio", False, True),
    (" (Toyota Corolla preto, placa RUP2I70), posteriormente recuperado em sua residência. "
     "O alvo da empreitada seria Felix; a ação, contudo, vitimou seu filho, Marcelo, em "
     "evidente erro na execução (aberratio ictus).", False, False),
])

p("Como adiante se demonstrará, a denúncia, no que toca ao respondente, é inepta e carece de "
  "justa causa: não descreve conduta concreta e individualizada que tenha sido indispensável "
  "ao resultado morte, e se ampara exclusivamente em prova indiciária frágil — “ouvir dizer”, "
  "delação de corréus contraditória e dados periféricos —, eivada de insuperáveis "
  "inconsistências fáticas, lógicas e geográficas.")

# ------------------------------------------------------------------ II - PRELIMINARES
heading("II — DAS PRELIMINARES")

heading("II.1 — Da inépcia da denúncia: ausência de descrição da conduta de José Raimundo indispensável ao evento morte (art. 41 c/c art. 395, I, do CPP)")

rich([
    ("Dispõe o art. 41 do CPP que a denúncia deve conter ", False, False),
    ("a exposição do fato criminoso com todas as suas circunstâncias", False, True),
    (". No tocante ao respondente, a inicial acusatória não cumpre tal ônus: limita-se a "
     "rotulá-lo, no fecho, como “mandante e coautor dos crimes”, sem narrar um único ato "
     "concreto de determinação, ajuste, pagamento ou direção da empreitada criminosa.", False, False),
])

p("Veja-se que a única passagem fática que o individualiza assim dispõe, ipsis litteris:")

quote("“JOSÉ RAIMUNDO SALES CHAVES JÚNIOR, vulgo ‘JÚNIOR BOLINHA’, mandante do homicídio de "
      "Felix da Silva Mendes Filho em virtude da disputa por uma máquina motoniveladora "
      "(patrol) (processo nº 0014566-68.2018.8.10.0001) que resultou, entretanto, na morte de "
      "Marcelo Mendes Martins, cedeu, inclusive, um dos veículos utilizados (Corolla, cor "
      "preta, placa RUP2I70), [que] foi, posteriormente, recuperado em sua residência.”")

p("A narrativa é meramente conclusiva. Afirma-se que o respondente é “mandante”, mas não se "
  "diz quando, onde, como, perante quem e em que termos ele teria contratado, ordenado ou "
  "remunerado os executores. Atribui-se a ele “contratar os serviços dos demais denunciados”, "
  "sem indicar o valor combinado, a forma de pagamento, o local do ajuste ou qualquer "
  "elemento concreto da suposta determinação.")

rich([
    ("Tratando-se de imputação a título de autoria mediata (mandante), era ", False, False),
    ("imprescindível", True, False),
    (" a descrição do nexo entre a alegada determinação e o resultado morte — o ato de "
     "comando sem o qual o crime não teria ocorrido. A denúncia, porém, é silente quanto a "
     "esse elo causal indispensável, o que inviabiliza o exercício da ampla defesa, pois não "
     "se defende de rótulos, mas de fatos.", False, False),
])

p("A imputação genérica, que não descreve a conduta individualizada do acusado, é repelida de "
  "forma pacífica pela jurisprudência dos Tribunais Superiores, por afronta ao art. 41 do CPP "
  "e às garantias do contraditório e da ampla defesa, impondo-se o reconhecimento da inépcia "
  "(art. 395, I, do CPP). Não por outra razão, idêntica preliminar de inépcia já foi suscitada "
  "nas respostas dos corréus Marcos Vinícius Campos e Luciano Rodrigues Ferreira.")

heading("II.2 — Da falta de justa causa: lastro probatório fundado em “ouvir dizer” e em delação de corréu sem corroboração (art. 395, III, c/c art. 155 do CPP)")

p("Ainda que superada a inépcia — o que se admite apenas para argumentar —, a ação penal, "
  "quanto ao respondente, carece de justa causa, porquanto desprovida do mínimo suporte "
  "probatório de autoria. Toda a imputação repousa em elementos indiretos e de nenhum valor "
  "incriminador autônomo, a saber:")

rich([
    ("a) ", True, False),
    ("“Ouvir dizer” (hearsay) da vítima sobrevivente. ", True, False),
    ("Felix afirmou ter ", False, False),
    ("ouvido dizer", False, True),
    (", por meio de conhecidos, que o corréu Marcos Vinícius estaria espalhando que o crime "
     "teria sido encomendado por “Júnior Bolinha”. Trata-se de testemunho indireto, sem "
     "qualquer valor probatório para sustentar a acusação, conforme firme orientação "
     "jurisprudencial.", False, False),
], first_line=True)

rich([
    ("b) ", True, False),
    ("Delação de corréu fundada em conversa alheia ouvida na cela. ", True, False),
    ("O corréu Marcos Vinícius — que busca isentar-se afirmando ter apenas emprestado seu "
     "veículo — declarou ter ", False, False),
    ("ouvido uma conversa", False, True),
    (" entre os corréus Gilbson e Luciano, no interior do presídio, sobre a suposta encomenda "
     "do crime por “Júnior Bolinha”. É delação de ouvir dizer sobre delação: prova de segundo "
     "grau, frágil e interessada, vedada como fundamento da persecução penal.", False, False),
], first_line=True)

rich([
    ("c) ", True, False),
    ("Delação isolada e contraditória do corréu Luciano (“Mix”). ", True, False),
    ("Afirmou ter ido cobrar uma dívida a mando do respondente — versão isolada, sem qualquer "
     "corroboração documental e frontalmente contraditada pelos demais elementos dos autos.", False, False),
], first_line=True)

rich([
    ("Como é cediço, ", False, False),
    ("a delação de corréu não se presta, isoladamente, a embasar nem a condenação nem a "
     "própria pretensão acusatória, exigindo corroboração por outros elementos de prova", True, False),
    (" — corroboração que, aqui, simplesmente inexiste. Soma-se a isso a vedação do art. 155 "
     "do CPP, que proíbe decisão fundada exclusivamente em elementos colhidos na fase "
     "inquisitorial. Ausente prova judicializável e idônea de autoria, falta justa causa "
     "(art. 395, III, do CPP).", False, False),
], first_line=True)

# ------------------------------------------------------------------ III - MÉRITO
heading("III — DAS INCONSISTÊNCIAS FÁTICAS, LÓGICAS E GEOGRÁFICAS")

p("Acaso ultrapassadas as preliminares, o exame do conjunto probatório revela contradições "
  "insanáveis que, desde já, anunciam a inevitável impronúncia do respondente ao término da "
  "primeira fase, por ausência de indícios suficientes de autoria (art. 414 do CPP).")

heading("III.1 — Inconsistência geográfica e temporal do rastreamento veicular: o veículo não estava sob a guarda de José Raimundo")

rich([
    ("A acusação sustenta que o respondente “cedeu” o Toyota Corolla preto (placa RUP2I70) "
     "para o crime. A própria prova técnica, contudo, desmente a tese. Os relatórios de "
     "rastreamento por GPS da locadora KINTO demonstram que o veículo ", False, False),
    ("pernoitou na residência do corréu Gilbson César (Rua das Acácias, Jardim Renascença) "
     "nas noites de 11 para 12 e de 12 para 13 de janeiro de 2023", True, False),
    (" — ou seja, o automóvel já estava sob o domínio e guarda do executor antes, durante e "
     "depois da data do crime (12/01/2023).", False, False),
])

p("Se o carro estava na posse de Gilbson, e nela permaneceu ao longo de todo o período dos "
  "fatos, não há como imputar a José Raimundo o ato de “ceder o veículo para o crime”. A "
  "ligação do respondente ao automóvel é estritamente post factum: ele só foi localizado em "
  "sua garagem em 25/01/2023, quase duas semanas depois do homicídio. Recuperação posterior "
  "de um veículo não é prova de participação no crime ocorrido semanas antes.")

heading("III.2 — A pessoa que entregou a chave do veículo não era José Raimundo")

rich([
    ("Reforça a fragilidade da prova material o fato de que, na recuperação de 25/01/2023, "
     "quem atendeu o preposto da locadora e lhe entregou a chave foi um homem ", False, False),
    ("de cerca de 60 anos de idade, compleição magra, cor parda, que não se identificou e que "
     "o próprio recuperador presumiu ser o caseiro da residência", True, False),
    (". O respondente, nascido em 1974, contava à época cerca de 48/49 anos. A descrição não "
     "corresponde à sua pessoa, e nenhuma testemunha o coloca, ele próprio, na posse ou "
     "entrega do veículo.", False, False),
])

heading("III.3 — Inconsistência geográfica da prova de telefonia (ERB): antena distante do local do crime e em horário incompatível")

rich([
    ("A acusação invoca o registro de uma ERB (antena) acionada por uma linha telefônica "
     "para situar o respondente. Tal elemento, porém, milita em favor da defesa:", False, False),
])

rich([
    ("(i) Titularidade alheia. ", True, False),
    ("A linha (98) 98423-1542 está cadastrada em nome da esposa do respondente, Betiane "
     "Meury; não há prova testemunhal ou visual de que fosse José Raimundo quem portava o "
     "aparelho. A vinculação a ele é mera dedução, extraída de uma chave PIX.", False, False),
], first_line=True)

rich([
    ("(ii) Distância do local do crime. ", True, False),
    ("A ERB foi acionada na Vila Embratel (lat. -2.564457; long. -44.312316), ao passo que o "
     "crime ocorreu no Rancho Félix, na Vila Maranhão/Maracanã (lat. -2.617936; long. "
     "-44.307504) — regiões distintas e distantes entre si. Em vez de aproximar, a prova "
     "afasta o aparelho do cenário do delito.", False, False),
], first_line=True)

rich([
    ("(iii) Horário incompatível. ", True, False),
    ("O acionamento ocorreu às 09h39 da manhã, ao passo que o crime se deu por volta das "
     "19h00 — mais de nove horas depois. Um sinal de antena pela manhã, em bairro diverso, "
     "nada prova quanto a uma participação em homicídio ocorrido à noite, em local distante.", False, False),
], first_line=True)

rich([
    ("(iv) Ligações de teor desconhecido. ", True, False),
    ("Os contatos telefônicos referidos ocorreram em 13/01/2023 (10h14 e 10h18), no dia "
     "seguinte ao crime, e ", False, False),
    ("seu teor é absolutamente desconhecido", False, True),
    (": colheu-se apenas a bilhetagem (registro de chamadas), jamais o conteúdo, pois não "
     "houve interceptação telefônica. Contato telefônico de conteúdo ignorado não demonstra "
     "ajuste criminoso.", False, False),
], first_line=True)

heading("III.4 — Ausência de reconhecimento e contradições entre os corréus quanto à motivação, arma e comando")

p("Nenhuma testemunha presencial reconheceu José Raimundo — formal ou informalmente. As "
  "vítimas e testemunhas oculares reconheceram exclusivamente os executores materiais "
  "(Luciano, Leilson e Carlos Augusto). O respondente jamais foi apontado por quem esteve no "
  "local dos fatos.")

p("A versão acusatória, ademais, é internamente contraditória:")

rich([
    ("• Quanto à motivação: ", True, False),
    ("o corréu Luciano fala em suposta dívida de R$ 23.000,00 em objetos de ouro (declaração "
     "oral, sem qualquer prova documental); a vítima Felix, por sua vez, refere desavença "
     "quanto a uma máquina motoniveladora. Versões inconciliáveis sobre o próprio móvel do "
     "crime.", False, False),
], first_line=True)

rich([
    ("• Quanto ao vínculo: ", True, False),
    ("o corréu Gilbson — apontado como o executor contratado — nega qualquer participação, "
     "nega o porte de arma e afirma conhecer José Raimundo apenas “da rotina da vida” (de "
     "frequentarem os mesmos bares), negando expressamente intimidade, negócio ou dívida com "
     "ele. A negativa do suposto contratado esvazia, pela base, a tese de contratação.", False, False),
], first_line=True)

rich([
    ("• Quanto à arma: ", True, False),
    ("a alegação de que o respondente teria fornecido a espingarda calibre 12 provém, "
     "novamente, de “conversa ouvida na cela” relatada por Marcos Vinícius, sem qualquer "
     "apreensão de arma em poder do respondente, perícia ou registro que o confirme.", False, False),
], first_line=True)

p("Não há, em suma: confissão; interceptação telefônica do respondente ordenando o crime; "
  "reconhecimento; prova de pagamento; prova de fornecimento de arma; nem prova de sua "
  "presença ou localização no sítio. Há, apenas, ilações construídas sobre “ouvir dizer”.")

heading("III.5 — Da vedação ao direito penal do autor: o indevido uso de fato pretérito (“caso Décio Sá”)")

rich([
    ("Os relatórios de investigação reiteradamente invocam o suposto envolvimento pretérito "
     "do respondente no caso do jornalista Décio Sá (processo nº 0006555-79.2020.8.10.0001) "
     "para “contextualizar sua periculosidade” e reforçar a imputação. Tal expediente é "
     "inadmissível: viola o ", False, False),
    ("direito penal do fato", True, False),
    (", segundo o qual ninguém pode ser processado por aquilo que supostamente é ou foi, mas "
     "tão somente por fato concreto, devidamente individualizado e provado. A invocação de "
     "fama e de processo diverso, em substituição à prova da autoria, denuncia exatamente a "
     "fragilidade do lastro probatório destes autos.", False, False),
])

# ------------------------------------------------------------------ IV - PROVAS
heading("IV — DA ESPECIFICAÇÃO DE PROVAS")

p("Protesta o respondente pela produção de todos os meios de prova em direito admitidos, em "
  "especial prova testemunhal, documental, pericial e o depoimento pessoal, bem como pela "
  "juntada oportuna de documentos, reservando-se o direito de, no curso da instrução, "
  "requerer diligências cuja necessidade se revele, notadamente:")

rich([("a) ", True, False),
      ("a oitiva do preposto/recuperador da locadora (Leonardo Costa Barros), a fim de "
       "esclarecer a identidade do homem que entregou a chave do veículo e a data/local de sua "
       "recuperação;", False, False)], first_line=True)
rich([("b) ", True, False),
      ("a requisição dos relatórios completos de rastreamento por GPS do veículo Corolla "
       "(locadora KINTO), comprovando sua permanência junto ao corréu Gilbson no período dos "
       "fatos;", False, False)], first_line=True)
rich([("c) ", True, False),
      ("a juntada integral dos laudos de ERB/bilhetagem, demonstrando a distância entre a "
       "antena acionada e o local do crime, bem como a ausência de interceptação de conteúdo;", False, False)], first_line=True)
rich([("d) ", True, False),
      ("a certidão de objeto e pé do processo nº 0014566-68.2018.8.10.0001, para esclarecer a "
       "natureza meramente patrimonial e cível da desavença quanto à motoniveladora.", False, False)], first_line=True)

p("Arrola, desde já, as testemunhas a seguir, requerendo suas intimações (rol a ser "
  "complementado no limite legal de 8 testemunhas, art. 406, §3º, do CPP):")

p("1. _______________________________________________________________;", first_line=False, space_after=2)
p("2. _______________________________________________________________;", first_line=False, space_after=2)
p("3. _______________________________________________________________.", first_line=False, space_after=10)

# ------------------------------------------------------------------ V - PEDIDOS
heading("V — DOS PEDIDOS")

p("Ante o exposto, requer:")

rich([("a) ", True, False),
      ("o recebimento e processamento da presente resposta, com a intimação da defesa de todos "
       "os atos do processo;", False, False)], first_line=True)

rich([("b) ", True, False),
      ("preliminarmente, o reconhecimento da ", False, False),
      ("inépcia da denúncia", True, False),
      (" quanto ao respondente (art. 395, I, c/c art. 41 do CPP), ante a ausência de descrição "
       "individualizada de conduta concreta indispensável ao evento morte;", False, False)], first_line=True)

rich([("c) ", True, False),
      ("sucessivamente, o reconhecimento da ", False, False),
      ("falta de justa causa", True, False),
      (" (art. 395, III, c/c art. 155 do CPP), por ausência de lastro probatório mínimo de "
       "autoria, fundado tão somente em “ouvir dizer” e em delação de corréu desprovida de "
       "corroboração;", False, False)], first_line=True)

rich([("d) ", True, False),
      ("no mérito, ao final da primeira fase do procedimento do júri, a ", False, False),
      ("IMPRONÚNCIA", True, False),
      (" do respondente, nos termos do art. 414 do CPP, ante a inexistência de indícios "
       "suficientes de autoria, em observância aos princípios da presunção de inocência e do "
       "in dubio pro reo (art. 5º, LVII, da CF);", False, False)], first_line=True)

rich([("e) ", True, False),
      ("a produção de todas as provas especificadas no item IV, com a intimação das "
       "testemunhas arroladas.", False, False)], first_line=True)

p("Termos em que,", first_line=False, space_after=2)
p("Pede deferimento.", first_line=False, space_after=18)

p("São Luís/MA, ____ de ______________ de 2026.", align='center', first_line=False, space_after=24)

p("_______________________________________", align='center', first_line=False, space_after=2)
p("ADVOGADO(A)", align='center', first_line=False, space_after=2)
p("OAB/MA nº __________", align='center', first_line=False, space_after=2)

doc.save(r"C:\Users\alden\.notebooklm\Resposta_Acusacao_Jose_Raimundo.docx")
print("Documento gerado com sucesso.")
