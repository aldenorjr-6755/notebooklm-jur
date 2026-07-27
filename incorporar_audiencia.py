# -*- coding: utf-8 -*-
import copy
from docx import Document
from docx.oxml.ns import qn

SRC = r"C:\Users\alden\OneDrive\Carlos Nina\Representacao_Etico_Disciplinar_OAB-MA_atualizada.docx"
OUT = r"C:\Users\alden\OneDrive\Carlos Nina\Representacao_Etico_Disciplinar_OAB-MA_v3.docx"

d = Document(SRC)

def set_single_text(p_el, text):
    runs = p_el.findall(qn('w:r'))
    if not runs:
        return
    first = runs[0]
    for r in runs[1:]:
        p_el.remove(r)
    ts = first.findall(qn('w:t'))
    if ts:
        ts[0].text = text
        ts[0].set(qn('xml:space'), 'preserve')
        for extra in ts[1:]:
            first.remove(extra)
    else:
        t = first.makeelement(qn('w:t'), {qn('xml:space'): 'preserve'})
        t.text = text
        first.append(t)

# ---- localizar o 1º fato (17/03/2021) e uma amostra de parágrafo "Normal" ----
ref_first = None
normal_sample = None
para_iv2 = None
ref_lastdoc = None
paras = d.paragraphs
for i, p in enumerate(paras):
    t = p.text.strip()
    if ref_first is None and t.startswith('17/03/2021') and p.style.name == 'List Paragraph':
        ref_first = p
        # amostra Normal = próximo parágrafo (linha "Processo ...")
        normal_sample = paras[i+1]
    if para_iv2 is None and t.startswith('Maria da Gl') and 'OAB/MA 6.399-A' in t:
        para_iv2 = p
    if '68685015' in t and p.style.name == 'List Paragraph':
        ref_lastdoc = p
assert ref_first is not None and normal_sample is not None, "não localizei o 1º fato"
assert ref_lastdoc is not None, "não localizei o último anexo"

# ===== novo fato 18/02/2021 (inserido ANTES do 1º fato) =====
header_txt = ("18/02/2021 – Ofensas orais em audiência: “advocacia espúria”, "
              "“maquiavelismo diabólico” e “stalking”")

corpo = [
    "Processo 0842458-45.2020 — audiência de medida protetiva de 18/02/2021 (registro audiovisual).",
    ("Autora das falas: Maria da Glória Costa Gonçalves de Sousa Aquino, em depoimento na "
     "audiência. O patrono Eduardo Corrêa, neste ato, não proferiu ofensas aos advogados."),
    ("(i) “...na petição do Dr. Carlos Nina, ele está negando que eu estava residindo "
     "em Portugal. Isso é uma falta de verdade...”;"),
    ("(ii) “E a gente nota um certo maquiavelismo diabólico nas petições, um discurso "
     "de ódio nas petições (...) do Dr. Carlos Nina.”;"),
    ("(iii) “É uma advocacia espúria (...) uma tentativa de manchar, de denegrir, que eu "
     "sei que não vai conseguir...”;"),
    ("(iv) “O Dr. Carlos Nina conturbou um pouco o processo (...) peticionando toda hora (...). Eu "
     "só não cedi a esse tipo de advocacia espúria...”;"),
    ("(v) “Quer dizer, é um tipo de advocacia que para mim seja brúculo, né?”;"),
    ("(vi) “Dr. Carlos Nina (...) o senhor veio fazendo um stalking na minha vida, inclusive "
     "fazendo levantamentos das movimentações processuais dos meus processos de "
     "inventário...”."),
    "Regra infringida: arts. 27 e 28 do CED (urbanidade e linguagem polida).",
]

anchor = ref_first._p
# 1) cabeçalho do fato (clone do item de lista numerada -> numera automaticamente como "1")
new_header = copy.deepcopy(ref_first._p)
set_single_text(new_header, header_txt)
anchor.addprevious(new_header)
# 2) corpo (clones de parágrafo Normal)
for txt in corpo:
    np = copy.deepcopy(normal_sample._p)
    set_single_text(np, txt)
    anchor.addprevious(np)
# 3) linha em branco separadora (clone de Normal vazio)
blank = copy.deepcopy(normal_sample._p)
set_single_text(blank, "")
anchor.addprevious(blank)

# ===== ajuste na individualização (IV.2 – Maria da Glória) =====
if para_iv2 is not None:
    runs = para_iv2._p.findall(qn('w:r'))
    if runs:
        last = runs[-1]
        ts = last.findall(qn('w:t'))
        add = (" Registre-se, ainda, que sua conduta ofensiva remonta à audiência de 18/02/2021, "
               "quando dirigiu oralmente ao Dr. Carlos Nina as expressões “advocacia espúria”, "
               "“maquiavelismo diabólico” e “stalking”.")
        if ts:
            ts[-1].text = (ts[-1].text or "") + add
            ts[-1].set(qn('xml:space'), 'preserve')

# ===== anexo: mídia da audiência =====
for t in ref_lastdoc._p.findall('.//' + qn('w:t')):
    if t.text and '68685015' in t.text:
        t.text = t.text.replace(').', ');')
new_anexo = copy.deepcopy(ref_lastdoc._p)
set_single_text(new_anexo, ("Mídia audiovisual da audiência de instrução de 18/02/2021, "
                            "no Processo 0842458-45.2020 (degravação anexa)."))
ref_lastdoc._p.addnext(new_anexo)

d.save(OUT)
print("Salvo em:", OUT)
print("Parágrafos:", len(Document(OUT).paragraphs))
